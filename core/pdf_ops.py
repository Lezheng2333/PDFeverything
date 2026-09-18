"""PDF 核心操作 — 所有 PDF 处理方法的统一入口。

所有耗时方法接受可选的 progress_callback(msg: str, pct: int) 参数，
供 GUI Worker 和 CLI 共享使用。"""

import os
import re
from pathlib import Path
from typing import Callable, List, Optional, Tuple

from .utils import check_input, ensure_output_dir

# PyMuPDF prints a one-time "consider pymupdf_layout" notice to stdout the first
# time table detection runs. On the MCP stdio channel that plain-text line breaks
# the JSON-RPC stream, and in the CLI it pollutes scriptable output — silence it
# once, before any MuPDF entry point is reached.
os.environ.setdefault("PYMUPDF_SUGGEST_LAYOUT_ANALYZER", "0")


def _silence_mupdf_notices() -> None:
    for mod_name in ("pymupdf", "fitz"):
        try:
            mod = __import__(mod_name)
            mod.no_recommend_layout()
            return
        except Exception:
            continue


_silence_mupdf_notices()

# Characters outside Latin-1 cannot be encoded by the base-14 PDF fonts
# (Helvetica/Courier). Rendering them through those fonts turns every glyph into
# a bullet, which silently destroys CJK content — so anything above U+00FF is
# drawn with a CJK-capable font instead.
_NON_LATIN = re.compile(r"[^\x00-\xff]")
_CJK_FONT_NAME = "china-s"
# find_tables() cost explodes with the number of vector segments on a page
# (measured: 600 lines ≈ 77 s). Skip detection above this budget.
TABLE_MAX_DRAWINGS = 200


def _needs_cjk(text: str) -> bool:
    return bool(_NON_LATIN.search(text or ""))


def _encrypted_page_count(path: Path) -> int:
    """Page count of a password-protected PDF (0 when it cannot be read).

    PyMuPDF reports the page count without needing the password, which is what
    `info` should show for a locked document."""
    try:
        import fitz
        doc = fitz.open(path)
        try:
            return len(doc)
        finally:
            doc.close()
    except Exception:
        return 0


def _prepare_output(path: Path) -> Path:
    """Make sure the output file's parent directory exists.

    Nearly every writer used to fail with a bare [Errno 2] when the user typed a
    path into a directory that did not exist yet."""
    ensure_output_dir(Path(path).parent)
    return Path(path)


class PdfOperator:
    """PDF 操作静态方法集合。CLI 和 GUI 共用。"""

    # ── 1. 查看信息 ─────────────────────────────────────

    @staticmethod
    def get_info(input_path: Path) -> dict:
        """返回 PDF 元信息字典。"""
        from pypdf import PdfReader

        check_input(input_path)
        size = input_path.stat().st_size

        # Reading metadata of an AES-encrypted PDF raises DependencyError when
        # `cryptography` is missing, and pypdf refuses to touch the pages of a
        # file it has not decrypted. Info must always work, so probe defensively
        # and fall back to PyMuPDF for the page count.
        encrypted = False
        pages = 0
        info = {}
        try:
            reader = PdfReader(input_path)
            encrypted = bool(reader.is_encrypted)
            if not encrypted:
                info = reader.metadata or {}
                pages = len(reader.pages)
            else:
                pages = _encrypted_page_count(input_path)
        except Exception:
            info = {}
            pages = _encrypted_page_count(input_path)

        return {
            "path": str(input_path),
            "pages": pages,
            "encrypted": encrypted,
            "size_bytes": size,
            "title": (info.title or "") if info else "",
            "author": (info.author or "") if info else "",
            "subject": (info.subject or "") if info else "",
            "creator": (info.creator or "") if info else "",
            "producer": (info.producer or "") if info else "",
        }

    # ── 2. 合并 ────────────────────────────────────────

    @staticmethod
    def merge(
        input_paths: List[Path],
        output_path: Path,
        progress_callback: Optional[Callable[[str, int], None]] = None,
    ) -> None:
        """将多个 PDF 合并为一个。"""
        from pypdf import PdfWriter

        if not input_paths:
            raise ValueError("没有输入文件")
        _prepare_output(output_path)
        writer = PdfWriter()
        total = len(input_paths)
        for i, p in enumerate(input_paths):
            check_input(p)
            writer.append(p)
            if progress_callback:
                progress_callback(f"合并中 ({i+1}/{total}): {p.name}", int((i+1)/total*100))

        writer.write(output_path)
        writer.close()

    # ── 3. 拆分 ────────────────────────────────────────

    @staticmethod
    def split(
        input_path: Path,
        output_dir: Path,
        page_ranges: Optional[List[Tuple[int, int]]] = None,
        progress_callback: Optional[Callable[[str, int], None]] = None,
    ) -> List[Path]:
        """将 PDF 拆分为单页或多页文件。

        page_ranges: [(1,5), (6,10)] 表示拆成两个文件（页码范围 1-based 包含）。
        如果为 None，则每页拆成一个文件。
        """
        from pypdf import PdfReader, PdfWriter

        check_input(input_path)
        reader = PdfReader(input_path)
        out_dir = ensure_output_dir(output_dir)
        stem = input_path.stem
        total_pages = len(reader.pages)

        if page_ranges is None:
            page_ranges = [(i, i) for i in range(1, total_pages + 1)]

        outputs = []
        total_ranges = len(page_ranges)
        for ri, (start, end) in enumerate(page_ranges):
            writer = PdfWriter()
            start_idx = max(0, start - 1)
            end_idx = min(total_pages, end)
            for pi in range(start_idx, end_idx):
                writer.add_page(reader.pages[pi])

            if end_idx <= start_idx:
                # A range that matches no page would otherwise produce an empty PDF.
                continue
            if end_idx - start_idx == 1:
                fname = f"{stem}_p{start:04d}.pdf"
            else:
                fname = f"{stem}_p{start:04d}-{end:04d}.pdf"
            out_path = out_dir / fname
            with open(out_path, "wb") as f:
                writer.write(f)
            outputs.append(out_path)

            if progress_callback:
                progress_callback(f"拆分中 ({ri+1}/{total_ranges})", int((ri+1)/total_ranges*100))

        return outputs

    # ── 4. 提取文字 ───────────────────────────────────

    @staticmethod
    def extract_text(
        input_path: Path,
        output_path: Path,
        progress_callback: Optional[Callable[[str, int], None]] = None,
    ) -> str:
        """提取 PDF 纯文本并写入文件（UTF-8）。返回完整文本。

        优先用 PyMuPDF（CJK 与版式还原都更好，且是核心依赖），失败时回退到
        pypdf，保证加密/损坏文件也有最后一次机会。"""
        check_input(input_path)
        _prepare_output(output_path)

        content = PdfOperator._extract_text_mupdf(input_path, progress_callback)
        if content is None:
            content = PdfOperator._extract_text_pypdf(input_path, progress_callback)

        output_path.write_text(content, encoding="utf-8")
        return content

    @staticmethod
    def _extract_text_mupdf(input_path: Path,
                            progress_callback=None) -> Optional[str]:
        """PyMuPDF text extraction. Returns None when the file cannot be read."""
        import fitz
        try:
            doc = fitz.open(input_path)
        except Exception:
            return None
        try:
            if doc.needs_pass:
                return None
            parts = []
            total = max(1, len(doc))
            for i, page in enumerate(doc, start=1):
                text = page.get_text("text", sort=True).strip()
                if text:
                    parts.append(text)
                if progress_callback:
                    progress_callback(f"提取文字 ({i}/{total})", int(i / total * 100))
            return "\n\n".join(parts)
        finally:
            doc.close()

    @staticmethod
    def _extract_text_pypdf(input_path: Path, progress_callback=None) -> str:
        from pypdf import PdfReader
        reader = PdfReader(input_path)
        parts = []
        total = max(1, len(reader.pages))
        for i, page in enumerate(reader.pages, start=1):
            text = (page.extract_text() or "").strip()
            if text:
                parts.append(text)
            if progress_callback:
                progress_callback(f"提取文字 ({i}/{total})", int(i / total * 100))
        return "\n\n".join(parts)

    # ── 5. 提取图片 ───────────────────────────────────

    @staticmethod
    def extract_images(
        input_path: Path,
        output_dir: Path,
        progress_callback: Optional[Callable[[str, int], None]] = None,
    ) -> int:
        """提取 PDF 中嵌入的图片。返回提取数量。"""
        import fitz  # PyMuPDF

        check_input(input_path)
        out_dir = ensure_output_dir(output_dir)
        doc = fitz.open(input_path)
        count = 0
        total = len(doc)
        for i, page in enumerate(doc, start=1):
            for j, img in enumerate(page.get_images(full=True), start=1):
                xref = img[0]
                try:
                    base_image = doc.extract_image(xref)
                except Exception:
                    continue  # broken/stencil/unsupported xref — skip this one
                ext = base_image.get("ext") or "png"
                img_path = out_dir / f"page{i:04d}_img{j:02d}.{ext}"
                img_path.write_bytes(base_image["image"])
                count += 1
            if progress_callback:
                progress_callback(f"提取图片 ({i}/{total})", int(i/total*100))
        doc.close()
        return count

    # ── 6. PDF → 图片 ─────────────────────────────────

    @staticmethod
    def to_images(
        input_path: Path,
        output_dir: Path,
        dpi: int = 200,
        progress_callback: Optional[Callable[[str, int], None]] = None,
    ) -> int:
        """将 PDF 每页转为 PNG 图片。返回页数。"""
        import fitz

        check_input(input_path)
        out_dir = ensure_output_dir(output_dir)
        stem = input_path.stem
        dpi = max(36, min(1200, int(dpi)))
        doc = fitz.open(input_path)
        total = len(doc)
        for i, page in enumerate(doc, start=1):
            pix = page.get_pixmap(dpi=dpi)
            img_path = out_dir / f"{stem}_p{i:04d}.png"
            pix.save(str(img_path))
            if progress_callback:
                progress_callback(f"转图片 ({i}/{total})", int(i/total*100))
        doc.close()
        return total

    # ── 7. 图片 → PDF ─────────────────────────────────

    @staticmethod
    def from_images(
        image_paths: List[Path],
        output_path: Path,
        progress_callback: Optional[Callable[[str, int], None]] = None,
    ) -> None:
        """将多张图片合并为一个 PDF。

        页面尺寸按图片自带的 DPI 换算为点（1 英寸 = 72 点），没有 DPI 信息时
        按 96 DPI 处理，避免手机照片变成一张 56 英寸宽的巨页。"""
        import fitz
        from PIL import Image as PILImage

        if not image_paths:
            raise ValueError("没有输入图片")
        _prepare_output(output_path)
        doc = fitz.open()
        total = len(image_paths)
        for i, p in enumerate(image_paths):
            check_input(p)
            with PILImage.open(p) as img:
                w, h = img.size
                dpi = img.info.get("dpi") or (96, 96)
            try:
                dpi_x = float(dpi[0]) or 96.0
                dpi_y = float(dpi[1]) or 96.0
            except (TypeError, ValueError, IndexError):
                dpi_x = dpi_y = 96.0
            pw, ph = w * 72.0 / dpi_x, h * 72.0 / dpi_y
            page = doc.new_page(width=pw, height=ph)
            page.insert_image(page.rect, filename=str(p))
            if progress_callback:
                progress_callback(f"加载图片 ({i+1}/{total}): {p.name}", int((i+1)/total*100))

        doc.save(output_path)
        doc.close()

    # ── 8. 压缩 ───────────────────────────────────────

    # 压缩档位 → (JPEG 质量, 目标 DPI)。None 表示不重采样图片。
    COMPRESS_MODES = {
        "lossless": (None, None),      # 只优化对象结构，画质完全不变
        "medium": (75, 150),           # 轻微降质，肉眼基本无感
        "max": (45, 96),               # 明显缩小，清晰度下降
    }

    @staticmethod
    def compress(
        input_path: Path,
        output_path: Path,
        mode: str = "lossless",
        progress_callback: Optional[Callable[[str, int], None]] = None,
    ) -> dict:
        """压缩 PDF。返回 {before_bytes, after_bytes, ratio, mode}。

        mode 决定压缩强度：
          lossless — 仅重写对象流（无损，画质不变）
          medium   — 图片重采样到 150 DPI / JPEG 质量 75
          max      — 图片重采样到 96 DPI / JPEG 质量 45
        """
        import pikepdf

        if mode not in PdfOperator.COMPRESS_MODES:
            raise ValueError(f"未知压缩档位: {mode}")
        quality, target_dpi = PdfOperator.COMPRESS_MODES[mode]

        check_input(input_path)
        _prepare_output(output_path)
        if progress_callback:
            progress_callback("正在压缩...", 20)

        if quality is None:
            pdf = pikepdf.open(input_path)
            try:
                pdf.save(
                    output_path,
                    compress_streams=True,
                    object_stream_mode=pikepdf.ObjectStreamMode.generate,
                )
            finally:
                pdf.close()
        else:
            PdfOperator._recompress_images(input_path, output_path, quality,
                                           target_dpi, progress_callback)

        before = input_path.stat().st_size
        after = output_path.stat().st_size
        ratio = (1 - after / before) * 100 if before > 0 else 0

        if progress_callback:
            progress_callback("压缩完成", 100)

        return {"before_bytes": before, "after_bytes": after,
                "ratio": ratio, "mode": mode}

    @staticmethod
    def _recompress_images(input_path: Path, output_path: Path, quality: int,
                           target_dpi: int, progress_callback=None) -> None:
        """Downsample and re-encode every large raster image, then rewrite the PDF.

        Images are replaced in place through the xref table, so page content,
        text, vectors and links are preserved exactly."""
        import io

        import fitz
        from PIL import Image as PILImage

        doc = fitz.open(input_path)
        try:
            seen = set()
            total = len(doc)
            for pno in range(total):
                page = doc[pno]
                if progress_callback:
                    progress_callback(f"压缩图片 ({pno+1}/{total})",
                                      int(20 + (pno + 1) / max(1, total) * 60))
                for img in page.get_images(full=True):
                    xref = img[0]
                    if xref in seen:
                        continue
                    seen.add(xref)
                    try:
                        info = doc.extract_image(xref)
                        raw = info["image"]
                        if len(raw) < 20_000:
                            continue  # too small to be worth re-encoding
                        with PILImage.open(io.BytesIO(raw)) as im:
                            im.load()
                            width, height = im.size
                            dpi = info.get("xres") or 0
                            # Only shrink when the image carries more pixels than
                            # the target DPI needs for its placement.
                            if dpi and dpi > target_dpi:
                                scale = target_dpi / dpi
                                width = max(1, int(width * scale))
                                height = max(1, int(height * scale))
                            if im.mode not in ("RGB", "L"):
                                im = im.convert("RGB")
                            if (width, height) != im.size:
                                im = im.resize((width, height), PILImage.LANCZOS)
                            buf = io.BytesIO()
                            im.save(buf, "JPEG", quality=quality, optimize=True)
                        new_bytes = buf.getvalue()
                        if len(new_bytes) < len(raw):
                            doc.update_stream(xref, new_bytes, compress=True)
                            doc.xref_set_key(xref, "Filter", "/DCTDecode")
                            doc.xref_set_key(xref, "ColorSpace", "/DeviceRGB")
                            doc.xref_set_key(xref, "Width", str(width))
                            doc.xref_set_key(xref, "Height", str(height))
                            doc.xref_set_key(xref, "BitsPerComponent", "8")
                    except Exception:
                        continue  # a single unreadable image must not fail the run
            doc.save(output_path, garbage=4, deflate=True, clean=True)
        finally:
            doc.close()

    # ── 9. 加水印 ─────────────────────────────────────

    @staticmethod
    def watermark(
        input_path: Path,
        watermark_path: Path,
        output_path: Path,
        progress_callback: Optional[Callable[[str, int], None]] = None,
    ) -> None:
        """给每一页叠加水印 PDF。"""
        from pypdf import PdfReader, PdfWriter

        check_input(input_path)
        check_input(watermark_path)
        _prepare_output(output_path)

        reader = PdfReader(input_path)
        watermark_reader = PdfReader(watermark_path)
        watermark_page = watermark_reader.pages[0]
        writer = PdfWriter()
        total = len(reader.pages)

        for i, page in enumerate(reader.pages):
            page.merge_page(watermark_page)
            writer.add_page(page)
            if progress_callback:
                progress_callback(f"加水印 ({i+1}/{total})", int((i+1)/total*100))

        with open(output_path, "wb") as f:
            writer.write(f)

    # ── 10. 加密 ──────────────────────────────────────

    @staticmethod
    def encrypt(
        input_path: Path,
        output_path: Path,
        password: str,
        progress_callback: Optional[Callable[[str, int], None]] = None,
    ) -> None:
        """给 PDF 设置打开密码。空密码会被拒绝——pypdf 会"加密"成不需要密码
        就能打开的文件，静默产生一个没有保护作用的 PDF。"""
        from pypdf import PdfReader, PdfWriter

        if not password or not password.strip():
            raise ValueError("密码不能为空")
        check_input(input_path)
        _prepare_output(output_path)
        if progress_callback:
            progress_callback("正在加密...", 50)

        reader = PdfReader(input_path)
        writer = PdfWriter()
        for page in reader.pages:
            writer.add_page(page)
        with open(output_path, "wb") as f:
            # AES-256 when the `cryptography` backend is available; the fallback
            # to RC4-128 must happen around BOTH the setup and the write, because
            # pypdf resolves the algorithm lazily and raises DependencyError while
            # serialising the file.
            try:
                writer.encrypt(password, algorithm="AES-256")
                writer.write(f)
            except Exception:
                if f.tell():
                    f.seek(0)
                    f.truncate()
                writer = PdfWriter()
                for page in reader.pages:
                    writer.add_page(page)
                writer.encrypt(password)
                writer.write(f)

        if progress_callback:
            progress_callback("加密完成", 100)

    # ── 11. 解密 ──────────────────────────────────────

    @staticmethod
    def decrypt(
        input_path: Path,
        output_path: Path,
        password: str,
        progress_callback: Optional[Callable[[str, int], None]] = None,
    ) -> None:
        """解密 PDF（移除密码）。密码错误时抛出 ValueError。"""
        from pypdf import PdfReader, PdfWriter

        check_input(input_path)
        _prepare_output(output_path)
        if progress_callback:
            progress_callback("正在解密...", 30)

        reader = PdfReader(input_path)
        if reader.is_encrypted:
            result = reader.decrypt(password)
            if result == 0:
                raise ValueError("密码不正确，无法解密")

        writer = PdfWriter()
        total = len(reader.pages)
        for i, page in enumerate(reader.pages):
            writer.add_page(page)
            if progress_callback:
                progress_callback(f"解密中 ({i+1}/{total})", 30 + int((i+1)/total*70))

        with open(output_path, "wb") as f:
            writer.write(f)

        if progress_callback:
            progress_callback("解密完成", 100)

    # ── 12. 旋转 ──────────────────────────────────────

    @staticmethod
    def rotate(
        input_path: Path,
        output_path: Path,
        angle: int,
        pages: Optional[List[int]] = None,
        progress_callback: Optional[Callable[[str, int], None]] = None,
    ) -> None:
        """旋转指定页面（90/180/270）。"""
        from pypdf import PdfReader, PdfWriter

        if angle not in (90, 180, 270):
            raise ValueError(f"旋转角度必须是 90/180/270，收到: {angle}")

        check_input(input_path)
        _prepare_output(output_path)
        reader = PdfReader(input_path)
        writer = PdfWriter()
        total = len(reader.pages)

        for i, page in enumerate(reader.pages):
            pnum = i + 1
            if pages is None or pnum in pages:
                page.rotate(angle)
            writer.add_page(page)
            if progress_callback:
                progress_callback(f"旋转中 ({pnum}/{total})", int(pnum/total*100))

        with open(output_path, "wb") as f:
            writer.write(f)

    # ── 13. 水印（文字） ─────────────────────────────

    @staticmethod
    def text_watermark(
        input_path: Path,
        output_path: Path,
        text: str,
        font_size: int = 60,
        opacity: float = 0.3,
        rotation: float = 45,
        progress_callback: Optional[Callable[[str, int], None]] = None,
    ) -> None:
        """在 PDF 每页上添加文字水印（对角线平铺）。

        rotation 为任意角度（用文本矩阵旋转，不再被圆整到 0/90/180/270）；
        opacity 真正生效（写入 ExtGState 的 /ca 与 /CA）；
        CJK 等非 Latin-1 字符会自动切换到内置中日韩字体，不再变成一个个圆点。
        """
        import fitz

        check_input(input_path)
        _prepare_output(output_path)
        opacity = max(0.0, min(1.0, float(opacity)))
        font = PdfOperator._watermark_font(text)

        doc = fitz.open(input_path)
        try:
            total = len(doc)
            for pi in range(total):
                page = doc[pi]
                rect = page.rect
                tw = fitz.TextWriter(rect, opacity=opacity, color=(0.7, 0.7, 0.7))
                for x_frac in (0.28, 0.62):
                    for y_frac in (0.22, 0.5, 0.78):
                        origin = fitz.Point(rect.x0 + rect.width * x_frac,
                                            rect.y0 + rect.height * y_frac)
                        tw.append(origin, text, font=font, fontsize=font_size)
                # Rotate the finished text around the page centre for a真·对角线水印.
                pivot = fitz.Point(rect.width / 2, rect.height / 2)
                morph = (pivot, fitz.Matrix(rotation))
                tw.write_text(page, morph=morph, overlay=True)
                if progress_callback:
                    progress_callback(f"添加文字水印 ({pi+1}/{total})",
                                      int((pi+1)/total*100))

            doc.save(output_path, incremental=False, garbage=3, deflate=True)
        finally:
            doc.close()

    @staticmethod
    def _watermark_font(text: str):
        """Pick a font that can actually encode `text`.

        Base-14 Helvetica covers Latin-1 only; anything beyond that renders as
        bullets, which is silent data corruption for Chinese/Japanese/Korean
        documents, so fall back to MuPDF's built-in CJK font."""
        import fitz
        if _needs_cjk(text):
            try:
                return fitz.Font(_CJK_FONT_NAME)
            except Exception:
                pass
        return fitz.Font("helv")

    # ── 14. PDF → Word ───────────────────────────────

    @staticmethod
    def to_word(
        input_path: Path,
        output_path: Path,
        progress_callback: Optional[Callable[[str, int], None]] = None,
    ) -> int:
        """Convert PDF to Word (.docx). Preserves text structure and tables.
        Returns number of pages processed."""
        import fitz
        from docx import Document
        from docx.shared import Pt, Inches

        import tempfile

        check_input(input_path)
        _prepare_output(output_path)
        doc_in = fitz.open(input_path)
        doc_out = Document()
        total = len(doc_in)

        try:
            for pi in range(total):
                page = doc_in[pi]
                if progress_callback:
                    progress_callback(f"PDF→Word ({pi+1}/{total})",
                                      int((pi+1)/total*100))

                doc_out.add_heading(f"Page {pi + 1}", level=1)

                # One dict extraction per page gives every span's font size; the
                # old code re-extracted the page for every block (O(n²)).
                span_sizes = PdfOperator._page_span_info(page)

                blocks = page.get_text("blocks")
                # PyMuPDF block tuples are (x0, y0, x1, y1, text, block_no,
                # block_type) — reading index 5 as block_type treated the second
                # text block as an image and discarded every later one.
                blocks.sort(key=lambda b: (b[1], b[0]))
                for x0, y0, x1, y1, text, block_no, block_type in blocks:
                    if block_type == 1:
                        PdfOperator._insert_block_image(
                            doc_out, page, (x0, y0, x1, y1), tempfile, Inches)
                        continue
                    if isinstance(text, bytes):
                        text = text.decode("utf-8", errors="replace")
                    text = text.strip()
                    if not text:
                        continue
                    if PdfOperator._block_font_size((x0, y0, x1, y1), span_sizes) >= 16:
                        doc_out.add_heading(text, level=2)
                    else:
                        doc_out.add_paragraph(text)

                for data in PdfOperator._page_tables(page):
                    if not data or not data[0]:
                        continue
                    ncols = len(data[0])
                    table = doc_out.add_table(rows=len(data), cols=ncols)
                    table.style = "Table Grid"
                    for ri, row_data in enumerate(data):
                        for ci, cell_val in enumerate(row_data):
                            if ci < ncols and cell_val is not None:
                                table.cell(ri, ci).text = str(cell_val)
        finally:
            doc_in.close()

        doc_out.save(str(output_path))
        return total

    # ── to_word / to_excel shared helpers ────────────

    @staticmethod
    def _page_span_info(page):
        """[(bbox, size)] for every text span on the page, in one pass."""
        spans = []
        try:
            raw = page.get_text("dict")
        except Exception:
            return spans
        for block in raw.get("blocks", []):
            if block.get("type") != 0:
                continue
            bbox = tuple(block.get("bbox", (0, 0, 0, 0)))
            size = 0.0
            for line in block.get("lines", []):
                for span in line.get("spans", []):
                    size = max(size, span.get("size", 0.0))
            if size:
                spans.append((bbox, float(size)))
        return spans

    @staticmethod
    def _block_font_size(bbox, span_sizes, default: float = 11.0) -> float:
        """Largest font size of the spans whose centre lies inside `bbox`."""
        x0, y0, x1, y1 = bbox
        best = 0.0
        for (bx0, by0, bx1, by1), size in span_sizes:
            cx, cy = (bx0 + bx1) / 2, (by0 + by1) / 2
            if x0 - 1 <= cx <= x1 + 1 and y0 - 1 <= cy <= y1 + 1:
                best = max(best, size)
        return best or default

    @staticmethod
    def _insert_block_image(doc_out, page, clip, tempfile_mod, inches_cls) -> None:
        """Render one image block to a temporary PNG and embed it in the docx."""
        tmp_path = None
        try:
            pix = page.get_pixmap(clip=clip, dpi=150)
            with tempfile_mod.NamedTemporaryFile(suffix=".png", delete=False) as tmp:
                tmp_path = tmp.name
                pix.save(tmp_path)
            doc_out.add_picture(tmp_path, width=inches_cls(5))
        except Exception:
            pass
        finally:
            if tmp_path:
                try:
                    os.unlink(tmp_path)
                except OSError:
                    pass

    @staticmethod
    def _page_tables(page) -> list:
        """Extract tables from a page as lists of rows.

        PyMuPDF 1.27 has no ``TableHeader.explicit_outer`` (the previous code
        raised AttributeError and silently produced no tables at all); and
        ``find_tables`` scales badly on drawing-heavy pages, so very busy pages
        are skipped rather than hanging for minutes."""
        try:
            if page.get_drawings().__len__() > TABLE_MAX_DRAWINGS:
                return []
        except Exception:
            pass
        try:
            found = page.find_tables()
        except Exception:
            return []
        out = []
        for tab in getattr(found, "tables", []) or []:
            try:
                data = tab.extract()
            except Exception:
                continue
            if data:
                out.append(data)
        return out

    # ── 15. PDF → PPT ────────────────────────────────

    @staticmethod
    def to_ppt(
        input_path: Path,
        output_path: Path,
        dpi: int = 200,
        progress_callback: Optional[Callable[[str, int], None]] = None,
    ) -> int:
        """Convert PDF to PowerPoint (.pptx). Each PDF page becomes one slide
        with the page rendered as a full-slide image."""
        import fitz
        from pptx import Presentation
        from pptx.util import Inches as PptInches

        import tempfile

        check_input(input_path)
        _prepare_output(output_path)
        dpi = max(36, min(1200, int(dpi)))
        doc = fitz.open(input_path)
        prs = Presentation()
        prs.slide_width = PptInches(13.333)   # widescreen
        prs.slide_height = PptInches(7.5)
        total = len(doc)

        try:
            with tempfile.TemporaryDirectory(prefix="pdfeverything_ppt_") as tmpdir:
                for pi in range(total):
                    page = doc[pi]
                    if progress_callback:
                        progress_callback(f"PDF→PPT ({pi+1}/{total})",
                                          int((pi+1)/total*100))
                    img_path = Path(tmpdir) / f"slide{pi+1:05d}.png"
                    page.get_pixmap(dpi=dpi).save(str(img_path))
                    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
                    slide.shapes.add_picture(
                        str(img_path), PptInches(0), PptInches(0),
                        width=prs.slide_width, height=prs.slide_height)
        finally:
            doc.close()

        prs.save(str(output_path))
        return total

    # ── 16. PDF → Excel ──────────────────────────────

    @staticmethod
    def to_excel(
        input_path: Path,
        output_path: Path,
        progress_callback: Optional[Callable[[str, int], None]] = None,
    ) -> int:
        """Convert PDF tables to Excel (.xlsx). Each table found in the PDF
        becomes a separate worksheet. Falls back to writing extracted text
        if no tables are detected."""
        import fitz
        from openpyxl import Workbook

        check_input(input_path)
        _prepare_output(output_path)
        doc = fitz.open(input_path)
        wb = Workbook()
        wb.remove(wb.active)   # drop the default sheet
        total = len(doc)
        sheet_count = 0

        try:
            for pi in range(total):
                page = doc[pi]
                if progress_callback:
                    progress_callback(f"PDF→Excel ({pi+1}/{total})",
                                      int((pi+1)/total*100))

                tables = PdfOperator._page_tables(page)
                for ti, data in enumerate(tables):
                    if not data:
                        continue
                    sheet_count += 1
                    ws = wb.create_sheet(title=PdfOperator._sheet_title(
                        f"Page{pi+1}_Tbl{ti+1}"))
                    for ri, row_data in enumerate(data, start=1):
                        for ci, cell_val in enumerate(row_data, start=1):
                            if cell_val is not None:
                                ws.cell(row=ri, column=ci, value=str(cell_val))

                # Text fallback — keyed on whether a table sheet was actually
                # produced, not on whether detection returned anything.
                if not tables:
                    text = page.get_text().strip()
                    if text:
                        sheet_count += 1
                        ws = wb.create_sheet(title=PdfOperator._sheet_title(
                            f"Page{pi+1}_Text"))
                        ws.cell(row=1, column=1,
                                value=f"Page {pi+1} — extracted text")
                        for li, line in enumerate(text.split("\n"), start=2):
                            ws.cell(row=li, column=1, value=line)
        finally:
            doc.close()

        if sheet_count == 0:
            ws = wb.create_sheet(title="Empty")
            ws.cell(row=1, column=1, value="No tables or text found in the PDF.")

        wb.save(str(output_path))
        return sheet_count

    @staticmethod
    def _sheet_title(name: str) -> str:
        """Excel sheet names: 31 chars max, no []:*?/\\ characters."""
        cleaned = re.sub(r"[\[\]:*?/\\]", "_", name)
        return cleaned[:31]
